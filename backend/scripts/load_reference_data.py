"""Load reference vocabulary into Supabase/Postgres tables.

Usage::

    python scripts/load_reference_data.py \\
        --manufacturers path/to/manufacturers.xlsx \\
        --lov path/to/lov.xlsx \\
        --uom path/to/uom.xlsx \\
        --fractions path/to/fractions.xlsx

When no ``.xlsx`` paths are supplied, derives a minimal vocabulary from a
Delivery-Format ground-truth CSV (``--fallback-csv``) instead.
"""
from __future__ import annotations

import argparse
import asyncio
import csv
import sys
import uuid
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from sqlalchemy import delete, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.db.models.reference import (
    FractionDecimal,
    LovAttribute,
    Manufacturer,
    UomAbbreviation,
)
from app.db.session import AsyncSessionLocal


def _cell_str(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _parse_sheet_rows(path: Path) -> list[list[str]]:
    """Read an xlsx sheet defensively (merged cells, multi-row headers)."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows: list[list[str]] = []
    for row in ws.iter_rows(values_only=True):
        cells = [_cell_str(c) for c in row]
        if any(cells):
            rows.append(cells)
    wb.close()
    return rows


def _find_header_row(rows: list[list[str]], markers: set[str]) -> int:
    for idx, row in enumerate(rows[:15]):
        lowered = {c.lower() for c in row if c}
        if markers & lowered:
            return idx
    return 0


def _column_index(header: list[str], *names: str) -> int | None:
    lowered = [h.lower() for h in header]
    for name in names:
        key = name.lower()
        if key in lowered:
            return lowered.index(key)
    return None


def load_manufacturers_xlsx(path: Path) -> list[dict]:
    rows = _parse_sheet_rows(path)
    header_idx = _find_header_row(
        rows, {"manufacturer_name", "manufacturer name", "mfr name"}
    )
    header = rows[header_idx]
    name_i = _column_index(header, "manufacturer_name", "manufacturer name", "mfr name")
    code_i = _column_index(header, "manufacturer_code", "manufacturer code", "mfr code")
    brand_i = _column_index(header, "brand_name", "brand name")
    brand_code_i = _column_index(header, "brand_code", "brand code")
    if name_i is None:
        raise ValueError(f"Could not locate manufacturer name column in {path}")

    records: list[dict] = []
    for row in rows[header_idx + 1 :]:
        name = row[name_i] if name_i < len(row) else ""
        if not name:
            continue
        records.append(
            {
                "manufacturer_name": name,
                "manufacturer_code": row[code_i] if code_i is not None and code_i < len(row) else "",
                "brand_name": row[brand_i] if brand_i is not None and brand_i < len(row) else "",
                "brand_code": row[brand_code_i] if brand_code_i is not None and brand_code_i < len(row) else "",
            }
        )
    return records


def load_lov_xlsx(path: Path) -> list[dict]:
    rows = _parse_sheet_rows(path)
    header_idx = _find_header_row(
        rows, {"classpath", "attribute_label", "attribute label"}
    )
    header = rows[header_idx]
    cp_i = _column_index(header, "classpath", "class path")
    leaf_i = _column_index(header, "leaf_node", "leaf node")
    filt_i = _column_index(header, "filtering")
    label_i = _column_index(header, "attribute_label", "attribute label")
    values_i = _column_index(header, "attribute_values", "attribute values")
    norm_label_i = _column_index(header, "normalized_label", "normalized label")
    norm_values_i = _column_index(header, "normalized_values", "normalized values")
    guide_i = _column_index(header, "guidelines")
    if label_i is None:
        raise ValueError(f"Could not locate attribute_label column in {path}")

    records: list[dict] = []
    for row in rows[header_idx + 1 :]:
        label = row[label_i] if label_i < len(row) else ""
        if not label:
            continue
        raw_values = row[values_i] if values_i is not None and values_i < len(row) else ""
        values = [v.strip() for v in raw_values.split(";") if v.strip()] if raw_values else []
        raw_norm = row[norm_values_i] if norm_values_i is not None and norm_values_i < len(row) else ""
        norm_values = [v.strip() for v in raw_norm.split(";") if v.strip()] if raw_norm else values
        records.append(
            {
                "classpath": row[cp_i] if cp_i is not None and cp_i < len(row) else "",
                "leaf_node": row[leaf_i] if leaf_i is not None and leaf_i < len(row) else "",
                "filtering": row[filt_i] if filt_i is not None and filt_i < len(row) else "",
                "attribute_label": label,
                "attribute_values": values,
                "normalized_label": row[norm_label_i] if norm_label_i is not None and norm_label_i < len(row) else label,
                "normalized_values": norm_values,
                "guidelines": row[guide_i] if guide_i is not None and guide_i < len(row) else "",
            }
        )
    return records


def load_uom_xlsx(path: Path) -> list[dict]:
    rows = _parse_sheet_rows(path)
    header_idx = _find_header_row(
        rows, {"measurement_type", "measurement type", "approved_abbreviation"}
    )
    header = rows[header_idx]
    type_i = _column_index(header, "measurement_type", "measurement type")
    abbr_i = _column_index(header, "approved_abbreviation", "approved abbreviation")
    var_i = _column_index(header, "accepted_variants", "accepted variants")
    ex_i = _column_index(header, "example")
    if abbr_i is None:
        raise ValueError(f"Could not locate approved_abbreviation column in {path}")

    records: list[dict] = []
    for row in rows[header_idx + 1 :]:
        abbr = row[abbr_i] if abbr_i < len(row) else ""
        if not abbr:
            continue
        raw_var = row[var_i] if var_i is not None and var_i < len(row) else ""
        variants = [v.strip() for v in raw_var.split(";") if v.strip()] if raw_var else []
        records.append(
            {
                "measurement_type": row[type_i] if type_i is not None and type_i < len(row) else "",
                "approved_abbreviation": abbr,
                "accepted_variants": variants,
                "example": row[ex_i] if ex_i is not None and ex_i < len(row) else "",
            }
        )
    return records


def load_fractions_xlsx(path: Path) -> list[dict]:
    rows = _parse_sheet_rows(path)
    header_idx = _find_header_row(rows, {"fraction", "decimal_value", "decimal"})
    header = rows[header_idx]
    frac_i = _column_index(header, "fraction")
    dec_i = _column_index(header, "decimal_value", "decimal", "decimal value")
    if frac_i is None or dec_i is None:
        raise ValueError(f"Could not locate fraction/decimal columns in {path}")

    records: list[dict] = []
    for row in rows[header_idx + 1 :]:
        frac = row[frac_i] if frac_i < len(row) else ""
        dec = row[dec_i] if dec_i < len(row) else ""
        if not frac or not dec:
            continue
        try:
            decimal_value = float(dec)
        except ValueError:
            continue
        records.append({"fraction": frac, "decimal_value": decimal_value})
    return records


def load_lov_from_delivery_csv(path: Path) -> list[dict]:
    """Fallback: walk ATTRIBUTE_* triplets from a Delivery-Format CSV."""
    records: list[dict] = []
    seen: set[tuple[str, str, str, str]] = set()
    with open(path, newline="", encoding="utf-8-sig") as handle:
        for record in csv.DictReader(handle):
            classpath = (record.get("Classpath") or "").strip()
            for i in range(1, 51):
                label = (record.get(f"ATTRIBUTE_LABEL {i}") or "").strip()
                value = (record.get(f"ATTRIBUTE_VALUE {i}") or "").strip()
                uom = (record.get(f"ATTRIBUTE_UOM {i}") or "").strip()
                if not label and not value:
                    continue
                key = (classpath, label, value, uom)
                if key in seen:
                    continue
                seen.add(key)
                display_value = value if not uom else f"{value} {uom}".strip()
                records.append(
                    {
                        "classpath": classpath,
                        "leaf_node": "",
                        "filtering": "",
                        "attribute_label": label,
                        "attribute_values": [display_value] if display_value else [],
                        "normalized_label": label.lower(),
                        "normalized_values": [display_value.lower()] if display_value else [],
                        "guidelines": "",
                    }
                )
    return records


async def _clear_and_insert(session: AsyncSession, model, rows: list[dict]) -> int:
    await session.execute(delete(model))
    for row in rows:
        session.add(model(id=uuid.uuid4(), **row))
    return len(rows)


async def _load(args: argparse.Namespace) -> None:
    has_xlsx = any(
        [args.manufacturers, args.lov, args.uom, args.fractions]
    )
    fallback = not has_xlsx

    if fallback:
        if not args.fallback_csv:
            print(
                "ERROR: No reference .xlsx files provided and no --fallback-csv given.",
                file=sys.stderr,
            )
            sys.exit(1)
        print(
            "\n*** WARNING: FALLBACK MODE ***\n"
            "Loading a minimum-viable vocabulary derived from the Delivery-Format CSV.\n"
            "This is NOT the full reference pack — demo LOV coverage will be limited.\n"
        )

    async with AsyncSessionLocal() as session:
        if fallback:
            lov_rows = load_lov_from_delivery_csv(Path(args.fallback_csv))
            count = await _clear_and_insert(session, LovAttribute, lov_rows)
            print(f"Inserted {count} lov_attributes rows (fallback).")
        else:
            if args.manufacturers:
                mfr_rows = load_manufacturers_xlsx(Path(args.manufacturers))
                count = await _clear_and_insert(session, Manufacturer, mfr_rows)
                print(f"Inserted {count} manufacturers rows.")
            if args.lov:
                lov_rows = load_lov_xlsx(Path(args.lov))
                count = await _clear_and_insert(session, LovAttribute, lov_rows)
                print(f"Inserted {count} lov_attributes rows.")
            if args.uom:
                uom_rows = load_uom_xlsx(Path(args.uom))
                count = await _clear_and_insert(session, UomAbbreviation, uom_rows)
                print(f"Inserted {count} uom_abbreviations rows.")
            if args.fractions:
                frac_rows = load_fractions_xlsx(Path(args.fractions))
                await session.execute(delete(FractionDecimal))
                for row in frac_rows:
                    session.add(FractionDecimal(**row))
                print(f"Inserted {len(frac_rows)} fraction_decimal rows.")

        await session.commit()
        result = await session.execute(select(LovAttribute).limit(1))
        if result.scalars().first() is None and fallback:
            print("No lov_attributes rows were inserted — check the fallback CSV.")


def main() -> None:
    parser = argparse.ArgumentParser(description="Load reference vocabulary tables.")
    parser.add_argument("--manufacturers", help="Path to manufacturers .xlsx")
    parser.add_argument("--lov", help="Path to LOV attributes .xlsx")
    parser.add_argument("--uom", help="Path to UOM abbreviations .xlsx")
    parser.add_argument("--fractions", help="Path to fraction/decimal .xlsx")
    parser.add_argument(
        "--fallback-csv",
        help="Delivery-Format CSV used when no .xlsx reference files are supplied",
    )
    args = parser.parse_args()
    asyncio.run(_load(args))


if __name__ == "__main__":
    main()
